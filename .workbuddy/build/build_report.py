"""Assemble the technical report.

Two passes are used. Pass 1 builds the document with placeholder page numbers in
the list of figures and the list of tables; Word then renders it and reports the
real page of every caption. Pass 2 rebuilds with those numbers and renders the
final PDF. The lists do not change length between passes, so pagination is stable.
"""
import json
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx.enum.text import WD_ALIGN_PARAGRAPH          # noqa: E402
from docx.shared import Pt                             # noqa: E402

from docxkit import MUTED, Report, add_field           # noqa: E402
import tech_a                                          # noqa: E402
import tech_b                                          # noqa: E402
import tech_c                                          # noqa: E402

OUT_DIR = os.environ.get("OUTDIR", r"C:\Users\Administrator\Desktop\JC2001-SF-Assessment\deliverables")
BASENAME = "JC2001_Group9_Technical_Report"


def collect():
    """Build a throwaway copy to learn the figure and table lists."""
    probe = Report(tech_a.TITLE)
    for fn in (tech_a.chapter1, tech_a.chapter2, tech_a.chapter3,
               tech_b.chapter4, tech_b.chapter5,
               tech_c.chapter6, tech_c.chapter7):
        fn(probe)
    return probe.figures, probe.tables


def page_numbers(path):
    """Read the caption -> page dump produced by Word, if present."""
    if not path or not os.path.exists(path):
        return {}
    out = {}
    with open(path, encoding="utf-8-sig") as f:
        for line in f:
            parts = line.strip().split("|")
            if len(parts) == 3:
                out["%s %s" % (parts[0], parts[1])] = parts[2]
    return out


def number_footer(R):
    s = R.doc.sections[0]
    s.different_first_page_header_footer = True
    p = s.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0
    add_field(p, " PAGE \\* ARABIC ", "2")
    for run in p.runs:
        run.font.size = Pt(10)
        run.font.name = "Times New Roman"


def build(pages=None):
    figures, tables = collect()
    pages = pages or {}

    def entry(label, caption):
        return (label, caption, pages.get(label, "—"))

    R = Report(tech_a.TITLE)
    number_footer(R)

    tech_a.title_page(R)
    R.toc("Table of Contents", depth=2)
    R.static_list("List of Figures", [entry(l, c) for l, c in figures])
    R.static_list("List of Tables", [entry(l, c) for l, c in tables])

    for fn in (tech_a.chapter1, tech_a.chapter2, tech_a.chapter3,
               tech_b.chapter4, tech_b.chapter5,
               tech_c.chapter6, tech_c.chapter7):
        fn(R)

    tech_c.references(R)
    tech_c.appendix_a(R)
    tech_c.appendix_b(R)
    tech_c.appendix_c(R)

    target = os.path.join(OUT_DIR, BASENAME + ".docx")
    R.save(target)
    print("report ->", target)
    print("  figures:", len(figures), " tables:", len(tables),
          " paragraphs:", len(R.doc.paragraphs))
    return target


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    kind = sys.argv[1] if len(sys.argv) > 1 else "1"
    pages = page_numbers(sys.argv[2]) if len(sys.argv) > 2 else None
    if kind == "collect":
        figures, tables = collect()
        print(json.dumps({"figures": figures, "tables": tables}, indent=1))
        return
    build(pages)


if __name__ == "__main__":
    main()
